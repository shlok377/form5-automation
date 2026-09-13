#!/usr/bin/env python3
"""
Excel (.xlsx) to CSV Ingestion Adapter for FORM-5
Converts uploaded or local Excel files into normalized, clean UTF-8 CSVs.
Handles:
- Datetime / Date formatting (DD/MM/YYYY)
- Float-to-integer normalization (e.g. 112590.0 -> 112590)
- Formula evaluation (reads calculated values)
- Empty/None cell normalization
- Multi-sheet selection (defaults to active sheet)
"""

import csv
import datetime
import io
import os
import sys

def format_cell_value(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:g}"
    if isinstance(val, bool):
        return "TRUE" if val else "FALSE"
    return str(val).strip()

def convert_xlsx_to_csv(xlsx_source, target_csv_path=None, sheet_name=None):
    """
    Converts an Excel (.xlsx) file or byte stream to normalized CSV format.
    
    Args:
        xlsx_source: File path (str), bytes, or file-like object.
        target_csv_path: Optional path to write the converted CSV file.
        sheet_name: Optional name of the sheet to convert (defaults to active sheet).
        
    Returns:
        The normalized CSV content as a UTF-8 string.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required for Excel support. Run: pip install openpyxl"
        )

    if isinstance(xlsx_source, (bytes, bytearray)):
        source_stream = io.BytesIO(xlsx_source)
        wb = openpyxl.load_workbook(source_stream, data_only=True)
    elif hasattr(xlsx_source, "read"):
        wb = openpyxl.load_workbook(xlsx_source, data_only=True)
    else:
        wb = openpyxl.load_workbook(str(xlsx_source), data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")

    for row in ws.iter_rows(values_only=True):
        # Format each cell
        formatted_row = [format_cell_value(cell) for cell in row]
        # Only write if row is not completely empty
        if any(c != "" for c in formatted_row):
            writer.writerow(formatted_row)

    csv_text = output.getvalue()

    if target_csv_path:
        os.makedirs(os.path.dirname(os.path.abspath(target_csv_path)), exist_ok=True)
        with open(target_csv_path, "w", encoding="utf-8-sig", newline="") as f:
            f.write(csv_text)

    return csv_text

def is_excel_file(filename):
    if not filename:
        return False
    fn = str(filename).lower().strip()
    return fn.endswith(".xlsx") or fn.endswith(".xlsm") or fn.endswith(".xltx")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        src = sys.argv[1]
        dst = sys.argv[2] if len(sys.argv) > 2 else None
        csv_out = convert_xlsx_to_csv(src, dst)
        if not dst:
            print(csv_out[:500])
        else:
            print(f"Successfully converted {src} to {dst}")
    else:
        print("Usage: python3 excel_adapter.py <input.xlsx> [output.csv]")
